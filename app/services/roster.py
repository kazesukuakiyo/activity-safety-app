"""名簿 (4 列の表) の取り込みと検証。

受け付ける入力:
  1. Excel (.xlsx) / CSV ファイル  … 大学指定テンプレート (学籍番号・氏名・所属・学年)
  2. 貼り付けテキスト              … Excel からコピーした 4 列 (タブ・カンマ区切り)
  3. 画面の行入力                  … 少人数の追加用

どの経路でも「4 列だけ」を強制し、それ以外の列 (要配慮情報など) は取り込まない。
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

ALLOWED_COLUMNS = ("学籍番号", "氏名", "所属", "学年")
_SPLIT = re.compile(r"\t|,|、|，")
MAX_ROWS = 500


@dataclass(frozen=True)
class RosterRow:
    student_no: str
    name: str
    department: str
    grade: str

    def as_line(self) -> str:
        return "\t".join((self.student_no, self.name, self.department, self.grade))


class RosterError(ValueError):
    pass


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel の数値セル (学籍番号 2024001.0 → 2024001)
    return str(value).strip()


def _rows_from_cells(table: list[list[str]], *, source: str = "") -> list[RosterRow]:
    """セルの 2 次元配列を検証して行にする。列数が 4 でない行はエラー。"""
    rows: list[RosterRow] = []
    errors: list[str] = []
    seen: set[str] = set()
    for lineno, cells in enumerate(table, start=1):
        cols = [_clean(c) for c in cells]
        while cols and cols[-1] == "":
            cols.pop()
        if not cols:
            continue
        if cols[:4] == list(ALLOWED_COLUMNS) and len(cols) == 4:
            continue  # 見出し行
        if len(cols) != 4:
            errors.append(f"{lineno} 行目: 列数が {len(cols)} です。学籍番号・氏名・所属・学年 の 4 列にしてください")
            continue
        if any(c == "" for c in cols):
            errors.append(f"{lineno} 行目: 空欄があります")
            continue
        row = RosterRow(*cols)
        if row.student_no in seen:
            errors.append(f"{lineno} 行目: 学籍番号 {row.student_no} が重複しています")
            continue
        seen.add(row.student_no)
        rows.append(row)
    if len(rows) > MAX_ROWS:
        errors.append(f"名簿は {MAX_ROWS} 人までです")
    if errors:
        raise RosterError("\n".join(errors))
    return rows


def parse_roster_text(text: str) -> list[RosterRow]:
    """貼り付けテキストを解析する。"""
    table = [_SPLIT.split(line.strip()) for line in (text or "").splitlines()]
    return _rows_from_cells(table)


def parse_roster_file(filename: str, data: bytes) -> list[RosterRow]:
    """Excel (.xlsx) または CSV ファイルを解析する。"""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise RosterError("Excel ファイルとして読み込めませんでした。テンプレートを使って .xlsx で保存してください") from None
        ws = wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return _rows_from_cells(table, source=filename)
    if name.endswith(".csv") or name.endswith(".txt") or name.endswith(".tsv"):
        text = None
        for enc in ("utf-8-sig", "cp932", "utf-8"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise RosterError("文字コードを判別できませんでした。UTF-8 か Shift_JIS で保存してください")
        first_line = text.splitlines()[0] if text.strip() else ""
        dialect = "excel-tab" if "\t" in first_line else "excel"
        table = [row for row in csv.reader(io.StringIO(text), dialect=dialect)]
        return _rows_from_cells(table, source=filename)
    if name.endswith(".xls"):
        raise RosterError("古い Excel 形式 (.xls) は読み込めません。「名前を付けて保存」で .xlsx にしてください")
    raise RosterError("対応しているのは Excel (.xlsx) と CSV です")


def rows_from_fields(student_nos, names, departments, grades) -> list[RosterRow]:
    """画面の行入力 (配列) から作る。すべて空の行は無視する。"""
    table = []
    for cells in zip(student_nos, names, departments, grades, strict=False):
        if any(_clean(c) for c in cells):
            table.append(list(cells))
    return _rows_from_cells(table)


def rows_to_text(rows) -> str:
    return "\n".join("\t".join((r.student_no, r.name, r.department, r.grade)) for r in rows)


def merge(*groups: list[RosterRow]) -> list[RosterRow]:
    """複数の入力元を学籍番号で重複除去して結合する。"""
    out: list[RosterRow] = []
    seen: set[str] = set()
    for g in groups:
        for r in g:
            if r.student_no not in seen:
                seen.add(r.student_no)
                out.append(r)
    return out


def to_csv(rows) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(ALLOWED_COLUMNS)
    for r in rows:
        w.writerow((r.student_no, r.name, r.department, r.grade))
    return "﻿" + buf.getvalue()  # BOM 付き (Excel で文字化けしない)


def build_template_xlsx(org_name: str = "", fiscal_year: int | None = None) -> bytes:
    """大学指定の名簿テンプレート (.xlsx) を生成する。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "部員名簿"
    ws.append(list(ALLOWED_COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F5")
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 8
    ws.freeze_panes = "A2"
    dv = DataValidation(
        type="list",
        formula1='"1,2,3,4,5,6,M1,M2,D1,D2,D3"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="学年",
        error="1〜6, M1, M2, D1〜D3 から選んでください",
    )
    ws.add_data_validation(dv)
    dv.add("D2:D501")
    # 学籍番号は文字列扱い (先頭 0 落ち防止)
    for r in range(2, 502):
        ws.cell(row=r, column=1).number_format = "@"

    note = wb.create_sheet("記入方法")
    lines = [
        "部員名簿 記入方法",
        "",
        f"団体: {org_name}" if org_name else "",
        f"年度: {fiscal_year}" if fiscal_year else "",
        "",
        "1. 「部員名簿」シートに 1 行 1 人で記入してください。",
        "2. 記入する項目は 学籍番号・氏名・所属・学年 の 4 つだけです。列を増やさないでください。",
        "3. 病歴、障害情報、常用薬、保護者連絡先などは記入しないでください (記入されたファイルは取り込めません)。",
        "4. 記入後、このファイルを .xlsx のまま保存し、システムの「年度部員名簿」からアップロードしてください。",
    ]
    for line in lines:
        note.append([line])
    note.column_dimensions["A"].width = 90
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
